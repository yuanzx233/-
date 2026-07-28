from __future__ import annotations

import json
import math
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / "tests" / "fixtures" / "dxf"
OUTPUT = ROOT / "output" / "tech-spike"

VALID_SITES = {
    "01_rectangle_18x24m": [(0, 0), (18000, 0), (18000, 24000), (0, 24000)],
    "02_rectangle_12x30m": [(0, 0), (12000, 0), (12000, 30000), (0, 30000)],
    "03_trapezoid": [(0, 0), (22000, 0), (18000, 26000), (2000, 26000)],
    "04_l_shape": [(0, 0), (24000, 0), (24000, 12000), (15000, 12000), (15000, 24000), (0, 24000)],
    "05_pentagon": [(0, 0), (18000, 0), (24000, 10000), (16000, 24000), (0, 20000)],
    "06_corner_lot": [(0, 0), (20000, 0), (20000, 20000), (0, 20000)],
    "07_compact_15x18m": [(0, 0), (15000, 0), (15000, 18000), (0, 18000)],
    "08_irregular_hexagon": [(0, 0), (16000, 0), (23000, 7000), (21000, 21000), (9000, 26000), (0, 18000)],
}


def dxf_pair(code: int, value: object) -> str:
    return f"{code}\n{value}\n"


def make_dxf(polygons: list[list[tuple[float, float]]], closed: bool = True, add_context: bool = True) -> str:
    parts = [dxf_pair(0, "SECTION"), dxf_pair(2, "HEADER"), dxf_pair(9, "$ACADVER"),
             dxf_pair(1, "AC1027"), dxf_pair(9, "$INSUNITS"), dxf_pair(70, 4),
             dxf_pair(0, "ENDSEC"), dxf_pair(0, "SECTION"), dxf_pair(2, "ENTITIES")]
    for points in polygons:
        parts += [dxf_pair(0, "LWPOLYLINE"), dxf_pair(8, "SITE_BOUNDARY"),
                  dxf_pair(90, len(points)), dxf_pair(70, 1 if closed else 0)]
        for x, y in points:
            parts += [dxf_pair(10, x), dxf_pair(20, y)]
    if add_context:
        parts += [dxf_pair(0, "LINE"), dxf_pair(8, "ROAD"), dxf_pair(10, -2000),
                  dxf_pair(20, -3000), dxf_pair(11, 26000), dxf_pair(21, -3000)]
        parts += [dxf_pair(0, "LINE"), dxf_pair(8, "ENTRANCE"), dxf_pair(10, 8500),
                  dxf_pair(20, 0), dxf_pair(11, 9500), dxf_pair(21, 0)]
        parts += [dxf_pair(0, "LINE"), dxf_pair(8, "NORTH"), dxf_pair(10, -1500),
                  dxf_pair(20, 0), dxf_pair(11, -1500), dxf_pair(21, 5000)]
    parts += [dxf_pair(0, "ENDSEC"), dxf_pair(0, "EOF")]
    return "".join(parts)


def generate_fixtures() -> None:
    FIXTURES.mkdir(parents=True, exist_ok=True)
    for name, points in VALID_SITES.items():
        (FIXTURES / f"{name}.dxf").write_text(make_dxf([points]), encoding="ascii")
    (FIXTURES / "invalid_open_boundary.dxf").write_text(
        make_dxf([[(0, 0), (18000, 0), (18000, 24000), (0, 24000)]], closed=False),
        encoding="ascii",
    )
    (FIXTURES / "invalid_multiple_boundaries.dxf").write_text(
        make_dxf([
            [(0, 0), (10000, 0), (10000, 10000), (0, 10000)],
            [(15000, 0), (25000, 0), (25000, 10000), (15000, 10000)],
        ]),
        encoding="ascii",
    )


def read_pairs(path: Path) -> list[tuple[int, str]]:
    lines = path.read_text(encoding="ascii", errors="strict").splitlines()
    if len(lines) % 2:
        raise ValueError("DXF_GROUP_PAIR_MISMATCH")
    return [(int(lines[i].strip()), lines[i + 1].strip()) for i in range(0, len(lines), 2)]


def parse_site_boundary(path: Path) -> list[tuple[float, float]]:
    pairs = read_pairs(path)
    boundaries: list[tuple[list[tuple[float, float]], bool]] = []
    i = 0
    while i < len(pairs):
        if pairs[i] == (0, "LWPOLYLINE"):
            i += 1
            layer, closed, points, x = "", False, [], None
            while i < len(pairs) and pairs[i][0] != 0:
                code, value = pairs[i]
                if code == 8:
                    layer = value
                elif code == 70:
                    closed = bool(int(value) & 1)
                elif code == 10:
                    x = float(value)
                elif code == 20 and x is not None:
                    points.append((x, float(value)))
                    x = None
                i += 1
            if layer == "SITE_BOUNDARY":
                boundaries.append((points, closed))
            continue
        i += 1
    if len(boundaries) != 1:
        raise ValueError("SITE_BOUNDARY_COUNT_INVALID")
    points, closed = boundaries[0]
    if not closed:
        raise ValueError("SITE_BOUNDARY_NOT_CLOSED")
    if len(points) < 3:
        raise ValueError("SITE_BOUNDARY_TOO_FEW_POINTS")
    return points


def metrics(points: list[tuple[float, float]]) -> dict[str, object]:
    area2 = sum(x1 * y2 - x2 * y1 for (x1, y1), (x2, y2) in zip(points, points[1:] + points[:1]))
    perimeter = sum(math.dist(a, b) for a, b in zip(points, points[1:] + points[:1]))
    xs, ys = zip(*points)
    return {
        "areaMm2": abs(area2) / 2,
        "perimeterMm": perimeter,
        "boundsMm": [min(xs), min(ys), max(xs), max(ys)],
        "vertexCount": len(points),
    }


def render_preview(points: list[tuple[float, float]], output: Path) -> None:
    image = Image.new("RGB", (960, 720), "#f3efe6")
    draw = ImageDraw.Draw(image)
    xs, ys = zip(*points)
    width, height = max(xs) - min(xs), max(ys) - min(ys)
    scale = min(760 / max(width, 1), 520 / max(height, 1))
    projected = [(100 + (x - min(xs)) * scale, 610 - (y - min(ys)) * scale) for x, y in points]
    draw.polygon(projected, fill="#d9e6d2", outline="#173d32", width=5)
    draw.line([(80, 640), (880, 640)], fill="#bd7048", width=24)
    draw.text((100, 40), f"Site preview | {abs(sum(x1*y2-x2*y1 for (x1,y1),(x2,y2) in zip(points, points[1:]+points[:1]))/2)/1_000_000:.1f} m2", fill="#173d32")
    image.save(output)


def write_obj(points: list[tuple[float, float]], output: Path, height: float = 7200) -> None:
    lines = ["# MVP site massing OBJ", "o HouseMassing"]
    for x, y in points:
        lines.append(f"v {x/1000:.3f} {y/1000:.3f} 0")
    for x, y in points:
        lines.append(f"v {x/1000:.3f} {y/1000:.3f} {height/1000:.3f}")
    n = len(points)
    lines.append("f " + " ".join(str(i) for i in range(1, n + 1)))
    lines.append("f " + " ".join(str(i) for i in range(2 * n, n, -1)))
    for i in range(n):
        a, b = i + 1, (i + 1) % n + 1
        lines.append(f"f {a} {b} {b+n} {a+n}")
    output.write_text("\n".join(lines) + "\n", encoding="ascii")


def write_pptx(preview: Path, output: Path, site_metrics: dict[str, object]) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.7), Inches(0.45), Inches(11.9), Inches(0.7))
    p = title.text_frame.paragraphs[0]
    p.text = "自建房概念方案 - 技术验证"
    p.font.size, p.font.bold = Pt(26), True
    slide.shapes.add_picture(str(preview), Inches(0.7), Inches(1.35), width=Inches(7.6))
    box = slide.shapes.add_textbox(Inches(8.7), Inches(1.6), Inches(4), Inches(3))
    tf = box.text_frame
    tf.text = f"地块面积\n{site_metrics['areaMm2']/1_000_000:.1f} m²"
    for text in [f"周长 {site_metrics['perimeterMm']/1000:.1f} m", "输出：PNG / OBJ / PPTX / PDF", "仅用于概念验证，不可直接施工"]:
        para = tf.add_paragraph()
        para.text = text
        para.space_before = Pt(14)
    prs.save(output)


def write_pdf(preview: Path, output: Path, site_metrics: dict[str, object]) -> None:
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    c = canvas.Canvas(str(output))
    c.setFont("STSong-Light", 20)
    c.drawString(50, 800, "自建房概念方案 - 技术验证")
    c.drawImage(str(preview), 50, 300, width=500, height=375, preserveAspectRatio=True)
    c.setFont("STSong-Light", 12)
    c.drawString(50, 270, f"地块面积：{site_metrics['areaMm2']/1_000_000:.1f} 平方米")
    c.drawString(50, 248, f"周长：{site_metrics['perimeterMm']/1000:.1f} 米")
    c.drawString(50, 90, "本成果为前期概念验证，不可直接用于施工、报建或结构安全判断。")
    c.save()


def run() -> dict[str, object]:
    generate_fixtures()
    OUTPUT.mkdir(parents=True, exist_ok=True)
    report: dict[str, object] = {"valid": {}, "invalid": {}, "artifacts": {}}
    for path in sorted(FIXTURES.glob("0*.dxf")):
        points = parse_site_boundary(path)
        report["valid"][path.name] = metrics(points)
    for path in sorted(FIXTURES.glob("invalid_*.dxf")):
        try:
            parse_site_boundary(path)
            report["invalid"][path.name] = "UNEXPECTED_SUCCESS"
        except ValueError as error:
            report["invalid"][path.name] = str(error)
    sample_path = FIXTURES / "01_rectangle_18x24m.dxf"
    sample_points = parse_site_boundary(sample_path)
    sample_metrics = metrics(sample_points)
    preview = OUTPUT / "site-preview.png"
    model = OUTPUT / "house-massing.obj"
    pptx = OUTPUT / "concept-package.pptx"
    pdf = OUTPUT / "concept-package.pdf"
    render_preview(sample_points, preview)
    write_obj(sample_points, model)
    write_pptx(preview, pptx, sample_metrics)
    write_pdf(preview, pdf, sample_metrics)
    report["artifacts"] = {k: str(v.relative_to(ROOT)) for k, v in {
        "preview": preview, "model": model, "pptx": pptx, "pdf": pdf
    }.items()}
    (OUTPUT / "validation-report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return report


if __name__ == "__main__":
    result = run()
    print(json.dumps(result, ensure_ascii=False, indent=2))
